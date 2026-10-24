"""Rebuild original Horizon demonstration media; requires Pillow, python-pptx,
reportlab, ffmpeg and espeak. Output is passed as the first argument.
Narration and transcript are generated from the same source, never inferred.
"""
import json, math, subprocess, sys, textwrap, wave
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from reportlab.pdfgen import canvas

OUT = Path(sys.argv[1] if len(sys.argv) > 1 else '/tmp/horizon-course')
OUT.mkdir(parents=True, exist_ok=True)
LESSONS = [
 ('embeddings', 'Meaning in numbers', [
  ('What is an embedding?', 'An embedding is a list of numbers that represents a piece of content. A trained model maps related meanings to nearby positions in a vector space. The numbers are not hand-written labels, and their meaning depends on the model.', 'Content → model → vector'),
  ('Similarity, not certainty', 'Imagine the phrases car repair and fixing an automobile. Their words differ, but their meanings are related. Embeddings can place them near each other. Cosine similarity measures how closely two vectors point in the same direction. It is a similarity score, not a probability of correctness.', 'car repair ≈ fixing an automobile'),
  ('A consistent vector space', 'Encode the query and the documents with the same compatible embedding model. Compare their vectors and rank candidate passages. An embedding does not verify a fact. It helps find material that might be useful, and the retrieved text still needs inspection.', 'Query vector → compare → ranked passages')
 ]),
 ('hybrid-search', 'Finding the right evidence', [
  ('Two ways to search', 'Keyword search rewards exact terms. Semantic search uses embeddings to match related meanings. An error code such as E forty two benefits from exact matching, while a paraphrased question may benefit from semantic matching. Neither method is best for every query.', 'Keyword precision + semantic meaning'),
  ('Combine ranked results', 'Hybrid retrieval combines keyword and vector search. Reciprocal rank fusion is one way to merge results. For each result, add one divided by a constant plus its rank in each list. This uses rank rather than pretending that different search scores share the same scale.', 'RRF score = Σ 1 / (k + rank)'),
  ('Permissions come first', 'Search only sources the learner is allowed to read. For a video question at the current position, exclude future transcript segments before ranking evidence. Preserve the source title and timestamp with each passage. Good relevance never overrides ownership or scope.', 'Allowed sources → eligible passages → ranking')
 ]),
 ('grounded-answers', 'From retrieval to grounded answers', [
  ('Evidence before answers', 'Retrieval augmented generation supplies selected source passages to a language model. The model can then draft an answer with references to those passages. A knowledge graph can connect concepts and resources, but a graph edge is not automatically proof that a claim is true.', 'Retrieve → inspect → answer → cite'),
  ('Citations must support claims', 'A useful citation identifies a source and its location, such as a video timestamp or a PDF page. Check that the cited passage actually supports the claim. Adding a citation label does not make an answer grounded, and retrieval does not eliminate hallucinations.', 'Claim + supporting passage + location'),
  ('Know when to abstain', 'If the permitted evidence is missing or contradictory, say that there is not enough support. Distinguish source-supported facts from general explanation. Grounded systems reduce risk through permission checks, scoped retrieval, source inspection, and honest uncertainty, not through a promise of perfect accuracy.', 'Insufficient evidence → explain the gap')
 ])]
serif = '/usr/share/fonts/noto/NotoSerif[wdth,wght].ttf'
sans = '/usr/share/fonts/noto/NotoSans[wdth,wght].ttf'
def font(size, display=False):
 return ImageFont.truetype(serif if display else sans, size)
metadata = []
for slug, title, sections in LESSONS:
 transcript, cursor, parts = [], 0.0, []
 for i, (heading, narration, formula) in enumerate(sections):
  wav = OUT / f'{slug}-{i}.wav'
  subprocess.run(['nix','run','nixpkgs#espeak','--','-s','145','-v','en-us','-w',str(wav),narration], check=True)
  with wave.open(str(wav)) as w: duration = w.getnframes()/w.getframerate()
  duration = math.ceil((duration + 1.0) * 25)/25
  transcript.append({'start': round(cursor,2), 'end': round(cursor+duration,2), 'text': narration, 'heading': heading})
  cursor += duration
  image = Image.new('RGB',(1280,720),'#f6f4ea'); d = ImageDraw.Draw(image)
  d.text((70,48),'HORIZON 2.0  /  DEMONSTRATION COURSE',font=font(18),fill='#a45232')
  d.text((70,120),heading,font=font(48,True),fill='#253b30')
  d.line((70,210,1210,210),fill='#d5d5c6',width=2)
  for j,line in enumerate(textwrap.wrap(narration, width=77)):
   d.text((70,245+j*37),line,font=font(26),fill='#343931')
  d.rounded_rectangle((70,570,1210,640),radius=8,fill='#253b30')
  d.text((100,588),formula,font=font(26),fill='#f6f4ea')
  d.text((70,670),f'{title}   /   {i+1} of {len(sections)}',font=font(18),fill='#6e7166')
  png = OUT / f'{slug}-{i}.png'; image.save(png)
  if i == 0: image.save(OUT / f'{slug}.jpg',quality=90)
  part = OUT / f'{slug}-{i}.mp4'; parts.append(part)
  subprocess.run(['ffmpeg','-v','error','-y','-loop','1','-i',str(png),'-i',str(wav),'-vf',"zoompan=z='min(zoom+0.00002,1.015)':d=1:s=1280x720:fps=25",'-af','apad','-t',str(duration),'-c:v','libx264','-preset','ultrafast','-crf','25','-pix_fmt','yuv420p','-c:a','aac',str(part)],check=True)
 listing = OUT / f'{slug}-concat.txt'; listing.write_text('\n'.join(f"file '{p}'" for p in parts))
 subprocess.run(['ffmpeg','-v','error','-y','-f','concat','-safe','0','-i',str(listing),'-c','copy','-movflags','+faststart',str(OUT/f'{slug}.mp4')],check=True)
 subprocess.run(['ffmpeg','-v','error','-y','-i',str(OUT/f'{slug}.mp4'),'-c:v','libvpx','-deadline','realtime','-cpu-used','8','-b:v','600k','-c:a','libopus',str(OUT/f'{slug}.webm')],check=True)
 metadata.append({'slug':slug,'title':title,'duration_seconds':round(cursor,2),'transcript':transcript})
 (OUT/f'{slug}.vtt').write_text('WEBVTT\n\n'+'\n\n'.join(f"{int(s['start']//3600):02}:{int(s['start']//60)%60:02}:{s['start']%60:06.3f} --> {int(s['end']//3600):02}:{int(s['end']//60)%60:02}:{s['end']%60:06.3f}\n{s['text']}" for s in transcript))
(OUT/'lessons.json').write_text(json.dumps(metadata,indent=2))
pdf = canvas.Canvas(str(OUT/'study-guide.pdf'),pagesize=(595,842))
presentation = Presentation(); presentation.slide_width=Inches(13.33); presentation.slide_height=Inches(7.5)
notes = ['# Foundations of grounded AI\n\nOriginal Horizon demonstration content. Narration uses synthetic speech.\n']
for slug,title,sections in LESSONS:
 pdf.setFont('Helvetica-Bold',22); pdf.drawString(45,785,title)
 pdf.setFont('Helvetica',10); pdf.drawString(45,760,'Horizon 2.0 | Original demonstration study guide')
 y=710
 for heading,narration,formula in sections:
  pdf.setFont('Helvetica-Bold',14); pdf.drawString(45,y,heading); y-=25
  pdf.setFont('Helvetica',11)
  for line in textwrap.wrap(narration,82): pdf.drawString(45,y,line); y-=17
  y-=30
  notes.append(f'## {heading}\n\n{narration}\n\n`{formula}`\n')
  slide=presentation.slides.add_slide(presentation.slide_layouts[1]); slide.shapes.title.text=heading
  body=slide.placeholders[1].text_frame; body.text=narration+'\n\n'+formula
  for p in body.paragraphs: p.font.size=Pt(24)
 pdf.showPage()
pdf.save(); presentation.save(OUT/'course-slides.pptx'); (OUT/'course-notes.md').write_text('\n'.join(notes))
(OUT/'hybrid-search.py').write_text('''"""Original demo: rank fusion, not an embedding service or AI model."""
def reciprocal_rank_fusion(keyword_ids, semantic_ids, k=60):
    scores = {}
    for ranking in (keyword_ids, semantic_ids):
        for rank, document_id in enumerate(dict.fromkeys(ranking), start=1):
            scores[document_id] = scores.get(document_id, 0) + 1 / (k + rank)
    return sorted(scores.items(), key=lambda item: item[1], reverse=True)

if __name__ == "__main__":
    print(reciprocal_rank_fusion(["exact-term", "shared"], ["shared", "paraphrase"]))
''')
print('Generated aligned lessons and resources at', OUT)